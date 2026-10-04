"""Safe text extraction for supported local document formats."""

from pathlib import Path
import zipfile


TEXT_EXTENSIONS = {".txt", ".md", ".rst"}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".pptx"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | DOCUMENT_EXTENSIONS
MAX_EXTRACTED_CHARACTERS = 500_000
MAX_OFFICE_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_OFFICE_MEMBERS = 2_000
MAX_COMPRESSION_RATIO = 100
MAX_PDF_PAGES = 200
MAX_SPREADSHEET_CELLS = 50_000
MAX_PRESENTATION_SLIDES = 200


class UnsupportedDocumentError(ValueError):
    pass


def extract_text(path: Path) -> str:
    extension = path.suffix.lower()
    if extension in TEXT_EXTENSIONS:
        return path.read_text(encoding="utf-8", errors="replace")[:MAX_EXTRACTED_CHARACTERS]
    if extension == ".docx":
        return _extract_docx(path)
    if extension == ".xlsx":
        return _extract_xlsx(path)
    if extension == ".pptx":
        return _extract_pptx(path)
    if extension == ".pdf":
        return _extract_pdf(path)
    raise UnsupportedDocumentError(f"Unsupported file type: {extension or '(no extension)'}")


def _extract_docx(path: Path) -> str:
    _validate_office_archive(path, "word/document.xml", "DOCX")

    try:
        from docx import Document
    except ImportError as error:
        raise RuntimeError("Install the document extra with `pip install -e '.[documents]'` for DOCX support.") from error

    document = Document(path)
    text_parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    for table in document.tables:
        for row in table.rows:
            text_parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(text_parts)[:MAX_EXTRACTED_CHARACTERS]


def _validate_office_archive(path: Path, required_member: str, format_name: str) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            uncompressed_size = sum(member.file_size for member in members)
            if len(members) > MAX_OFFICE_MEMBERS or uncompressed_size > MAX_OFFICE_UNCOMPRESSED_BYTES:
                raise UnsupportedDocumentError(f"{format_name} archive exceeds extraction safety limits.")
            if any(
                member.file_size > 0
                and member.file_size / max(member.compress_size, 1) > MAX_COMPRESSION_RATIO
                for member in members
            ):
                raise UnsupportedDocumentError(f"{format_name} archive compression ratio exceeds safety limits.")
            if required_member not in archive.namelist():
                raise UnsupportedDocumentError(f"{format_name} archive does not contain its document body.")
    except zipfile.BadZipFile as error:
        raise UnsupportedDocumentError(f"{format_name} file is not a valid ZIP archive.") from error


def _extract_xlsx(path: Path) -> str:
    _validate_office_archive(path, "xl/workbook.xml", "XLSX")
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("Install the document extra with `pip install -e '.[documents]'` for XLSX support.") from error

    workbook = load_workbook(path, read_only=True, data_only=True)
    text_parts = []
    extracted_characters = 0
    visited_cells = 0
    try:
        for worksheet in workbook.worksheets:
            if extracted_characters >= MAX_EXTRACTED_CHARACTERS:
                break
            text_parts.append(f"Sheet: {worksheet.title}")
            extracted_characters += len(text_parts[-1])
            for row in worksheet.iter_rows(values_only=True):
                if visited_cells >= MAX_SPREADSHEET_CELLS or extracted_characters >= MAX_EXTRACTED_CHARACTERS:
                    break
                visited_cells += len(row)
                values = [str(value) for value in row if value is not None]
                if not values:
                    continue
                row_text = "\t".join(values)
                remaining = MAX_EXTRACTED_CHARACTERS - extracted_characters
                text_parts.append(row_text[:remaining])
                extracted_characters += min(len(row_text), remaining)
    finally:
        workbook.close()
    return "\n".join(text_parts)


def _extract_pptx(path: Path) -> str:
    _validate_office_archive(path, "ppt/presentation.xml", "PPTX")
    try:
        from pptx import Presentation
    except ImportError as error:
        raise RuntimeError("Install the document extra with `pip install -e '.[documents]'` for PPTX support.") from error

    try:
        presentation = Presentation(path)
        if len(presentation.slides) > MAX_PRESENTATION_SLIDES:
            raise UnsupportedDocumentError("PPTX slide count exceeds the extraction safety limit.")
        text_parts = []
        extracted_characters = 0
        for slide_number, slide in enumerate(presentation.slides, start=1):
            if extracted_characters >= MAX_EXTRACTED_CHARACTERS:
                break
            slide_title = f"Slide {slide_number}"
            text_parts.append(slide_title)
            extracted_characters += len(slide_title)
            for shape in slide.shapes:
                if shape.has_table:
                    shape_text = "\n".join(
                        "\t".join(cell.text for cell in row.cells)
                        for row in shape.table.rows
                    )
                elif shape.has_text_frame:
                    shape_text = shape.text
                else:
                    continue
                remaining = MAX_EXTRACTED_CHARACTERS - extracted_characters
                if remaining <= 0:
                    break
                text_parts.append(shape_text[:remaining])
                extracted_characters += min(len(shape_text), remaining)
        return "\n".join(text_parts)
    except UnsupportedDocumentError:
        raise
    except Exception as error:
        raise UnsupportedDocumentError("PPTX text extraction failed.") from error


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as error:
        raise RuntimeError("Install the document extra with `pip install -e '.[documents]'` for PDF support.") from error

    try:
        reader = PdfReader(path, strict=False)
        if reader.is_encrypted:
            raise UnsupportedDocumentError("Encrypted PDF files cannot be indexed.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise UnsupportedDocumentError("PDF page count exceeds the extraction safety limit.")
        text_parts = []
        extracted_characters = 0
        for page in reader.pages:
            remaining = MAX_EXTRACTED_CHARACTERS - extracted_characters
            if remaining <= 0:
                break
            page_text = page.extract_text() or ""
            page_text = page_text[:remaining]
            text_parts.append(page_text)
            extracted_characters += len(page_text)
        return "\n".join(text_parts)
    except UnsupportedDocumentError:
        raise
    except Exception as error:
        raise UnsupportedDocumentError("PDF text extraction failed.") from error