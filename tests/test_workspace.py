"""Application behavior tests: real database, graph, extraction, and HTTP routes."""

import json
import time
import zipfile
from io import BytesIO

import pytest
from fastapi.testclient import TestClient

from file_recommender.api import create_app
from file_recommender.index import IndexStore
from file_recommender.worker import Worker


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    for key in (
        "FILE_RECOMMENDER_AUTH_TOKENS",
        "FILE_RECOMMENDER_OIDC_ISSUER",
        "FILE_RECOMMENDER_OIDC_AUDIENCE",
        "FILE_RECOMMENDER_OIDC_JWKS_URL",
    ):
        monkeypatch.delenv(key, raising=False)
    store = IndexStore(tmp_path / "data/index.sqlite3")
    app = create_app(store)
    with TestClient(app) as client:
        token = client.get("/api/v1/bootstrap").json()["token"]
        client.headers["X-Local-Token"] = token
        yield client, app.state.library, Worker(app.state.library)


def connect(client, worker, folder):
    response = client.post("/api/v1/sources", json={"path": str(folder)})
    assert response.status_code == 201, response.text
    assert worker.once()
    job = client.get("/api/v1/jobs/" + response.json()["job_id"]).json()
    assert job["status"] == "completed", job
    return response.json()["source"]["id"]


def search(client, query, session=None, filters=None):
    session = session or client.post("/api/v1/sessions").json()["id"]
    response = client.post(
        f"/api/v1/sessions/{session}/turns", json={"query": query, "filters": filters or {}}
    )
    assert response.status_code == 202, response.text
    turn_id = response.json()["request_id"]
    for _ in range(200):
        turn = client.get(f"/api/v1/turns/{turn_id}").json()
        if turn["status"] not in {"queued", "running"}:
            assert turn["status"] == "completed", turn
            return session, turn["result"]
        time.sleep(0.02)
    pytest.fail("Search did not terminate.")


def test_folder_search_followup_preview_feedback_and_download(workspace, tmp_path):
    client, library, worker = workspace
    folder = tmp_path / "documents"
    folder.mkdir()
    (folder / "orchid-plan.md").write_text(
        "# Orchid migration\nPlatform rollout milestones and delivery timeline."
    )
    (folder / "camera-notes.txt").write_text("Camera lens aperture and photography lessons.")
    source_id = connect(client, worker, folder)
    _, result = search(client, "orchid migration", filters={"source_id": source_id})
    assert result["results"][0]["name"] == "orchid-plan.md"
    assert result["diagnostics"]["candidate_count"] == 1
    item = result["results"][0]
    preview = client.get(f"/api/v1/files/{item['id']}/preview").json()
    assert preview["sections"][0]["label"] == "# Orchid migration"
    assert client.get(f"/api/v1/files/{item['id']}/download").content.startswith(b"# Orchid")
    assert (
        client.post(
            "/api/v1/feedback",
            json={"recommendation_id": item["recommendation_id"], "feedback": "relevant"},
        ).status_code
        == 200
    )
    session, result = search(client, "migration")
    _, result = search(client, "only PDFs", session)
    assert result["filters"]["extension"] == "pdf"
    assert result["results"] == []
    assert "migration" in result["effective_query"]
    assert client.get("/api/v1/health").json()["file_count"] == 2


def test_incremental_preserves_ids_and_failed_files_and_unavailable_root(workspace, tmp_path):
    client, library, worker = workspace
    folder = tmp_path / "docs"
    folder.mkdir()
    file = folder / "notes.txt"
    file.write_text("first project plan")
    source_id = connect(client, worker, folder)
    original = client.get("/api/v1/files").json()["files"][0]["id"]
    client.post(f"/api/v1/sources/{source_id}/sync")
    worker.once()
    assert client.get("/api/v1/jobs").json()["jobs"][0]["indexed"] == 0
    file.write_text("revised project milestones")
    client.post(f"/api/v1/sources/{source_id}/sync")
    worker.once()
    assert client.get("/api/v1/files").json()["files"][0]["id"] == original
    assert client.get(f"/api/v1/files/{original}").json()["version"] == 2
    file.write_text("")
    client.post(f"/api/v1/sources/{source_id}/sync")
    worker.once()
    assert client.get("/api/v1/health").json()["file_count"] == 1
    folder.rename(tmp_path / "gone")
    client.post(f"/api/v1/sources/{source_id}/sync")
    worker.once()
    assert client.get("/api/v1/jobs").json()["jobs"][0]["status"] == "failed"
    assert client.get("/api/v1/health").json()["file_count"] == 1


def test_sources_do_not_overlap_and_removal_keeps_originals(workspace, tmp_path):
    client, _, worker = workspace
    folder = tmp_path / "docs"
    folder.mkdir()
    (folder / "file.txt").write_text("project details")
    source_id = connect(client, worker, folder)
    assert client.post("/api/v1/sources", json={"path": str(folder)}).status_code == 400
    assert client.post("/api/v1/sources", json={"path": str(tmp_path)}).status_code == 400
    assert client.delete(f"/api/v1/sources/{source_id}").status_code == 200
    assert (folder / "file.txt").is_file()
    assert client.get("/api/v1/health").json()["file_count"] == 0


def test_uploads_use_safe_names_and_backup_restore(workspace):
    client, _, worker = workspace
    response = client.post(
        "/api/v1/uploads",
        files=[
            ("files", ("../../plan.md", b"Migration launch timeline", "text/plain")),
            ("files", ("bad.exe", b"abc", "application/octet-stream")),
        ],
    )
    assert response.status_code == 202, response.text
    assert response.json()["accepted"] == ["plan.md"]
    assert response.json()["rejected"][0]["file"] == "bad.exe"
    worker.once()
    item = client.get("/api/v1/files").json()["files"][0]
    assert item["name"] == "plan.md" and ".." not in item["path"]
    backup = client.get("/api/v1/backup")
    assert backup.status_code == 200
    with zipfile.ZipFile(BytesIO(backup.content)) as z:
        assert "index.sqlite3" in z.namelist()
    assert client.delete("/api/v1/sources/uploads").status_code == 200
    assert client.get("/api/v1/health").json()["file_count"] == 0
    restored = client.post(
        "/api/v1/restore", files={"archive": ("backup.zip", backup.content, "application/zip")}
    )
    assert restored.status_code == 200, restored.text
    assert client.get("/api/v1/health").json()["file_count"] == 1
    assert (
        client.get(f"/api/v1/files/{item['id']}/download").content == b"Migration launch timeline"
    )


def test_csrf_origin_host_and_symlink_download_boundary(workspace, tmp_path):
    client, _, worker = workspace
    assert (
        client.get("/api/v1/sources", headers={"Origin": "https://evil.example"}).status_code == 403
    )
    assert client.get("/api/v1/bootstrap", headers={"Host": "evil.example"}).status_code == 400
    assert client.post("/api/v1/sessions", headers={"X-Local-Token": "wrong"}).status_code == 403
    folder = tmp_path / "docs"
    folder.mkdir()
    path = folder / "notes.txt"
    path.write_text("original contents")
    connect(client, worker, folder)
    file_id = client.get("/api/v1/files").json()["files"][0]["id"]
    secret = tmp_path / "secret.txt"
    secret.write_text("private")
    path.unlink()
    path.symlink_to(secret)
    assert client.get(f"/api/v1/files/{file_id}/download").status_code == 404


def test_cancel_queued_job_does_not_modify_index(workspace, tmp_path):
    client, _, worker = workspace
    folder = tmp_path / "docs"
    folder.mkdir()
    (folder / "file.txt").write_text("index me")
    response = client.post("/api/v1/sources", json={"path": str(folder)})
    job_id = response.json()["job_id"]
    assert client.post(f"/api/v1/jobs/{job_id}/cancel").status_code == 200
    assert not worker.once()
    assert client.get("/api/v1/health").json()["file_count"] == 0


def test_source_scope_and_filter_survive_expansion_and_model_failure(
    workspace, tmp_path, monkeypatch
):
    client, library, worker = workspace
    for name, text in (
        ("one", "Automobile servicing vehicle repair"),
        ("two", "Automobile servicing vehicle repair private"),
    ):
        folder = tmp_path / name
        folder.mkdir()
        (folder / "guide.txt").write_text(text)
        connect(client, worker, folder)
    selected = client.get("/api/v1/sources").json()["sources"][0]["id"]
    library.set_settings({"llm_provider": "ollama"})
    from file_recommender.agent import Agent

    def fail(*args):
        raise TimeoutError()

    monkeypatch.setattr(Agent, "analyze", fail)
    _, result = search(
        client,
        "find automobile servicing maintenance repair instructions",
        filters={"source_id": selected, "extension": "txt"},
    )
    assert all(item["source_id"] == selected for item in result["results"])
    assert result["filters"]["extension"] == "txt"
    assert result["diagnostics"]["fallback"]


def test_personalization_reset_and_secret_settings(workspace, tmp_path):
    client, library, worker = workspace
    folder = tmp_path / "docs"
    folder.mkdir()
    (folder / "notes.txt").write_text("project notes planning")
    connect(client, worker, folder)
    file_id = client.get("/api/v1/files").json()["files"][0]["id"]
    client.get(f"/api/v1/files/{file_id}/preview")
    with library.connect() as db:
        assert db.execute("SELECT COUNT(*) FROM access_events").fetchone()[0] == 1
    assert client.post("/api/v1/settings/reset-activity").status_code == 200
    with library.connect() as db:
        assert db.execute("SELECT COUNT(*) FROM access_events").fetchone()[0] == 0
    assert (
        client.patch(
            "/api/v1/settings", json={"cloud_url": "http://insecure.example/v1"}
        ).status_code
        == 400
    )
    assert client.patch("/api/v1/settings", json={"cloud_key": "secret-key"}).status_code == 200
    settings = client.get("/api/v1/settings").json()
    assert "cloud_key" not in settings and settings["cloud_key_set"]
    assert client.patch("/api/v1/settings", json={"personalization": None}).status_code == 400


def test_unsafe_restore_is_rejected_without_changing_index(workspace):
    client, _, _ = workspace
    payload = BytesIO()
    with zipfile.ZipFile(payload, "w") as z:
        z.writestr("../escape.txt", "malicious")
    response = client.post(
        "/api/v1/restore", files={"archive": ("bad.zip", payload.getvalue(), "application/zip")}
    )
    assert response.status_code == 400
    assert client.get("/api/v1/health").json()["status"] == "ready"


def test_legacy_auth_does_not_expose_workspace(tmp_path, monkeypatch):
    monkeypatch.setenv("FILE_RECOMMENDER_AUTH_TOKENS", json.dumps({"alice": "a" * 32}))
    with TestClient(create_app(IndexStore(tmp_path / "data/index.sqlite3"))) as client:
        assert client.get("/api/v1/bootstrap").status_code == 503


def test_native_vector_search_scopes_sources_and_model_versions(workspace, tmp_path, monkeypatch):
    import struct

    from file_recommender.agent import Agent

    client, library, worker = workspace
    source_ids = []
    for name in ("one", "two"):
        folder = tmp_path / name
        folder.mkdir()
        (folder / "notes.txt").write_text("Technical engineering architecture")
        source_ids.append(connect(client, worker, folder))
    with library.connect() as db:
        for row in db.execute("SELECT id,content_hash FROM file_chunks").fetchall():
            db.execute(
                "INSERT INTO chunk_embeddings VALUES(?,?,'fixture-v1',2,?)",
                (row[0], row[1], struct.pack("<2f", 1, 0)),
            )
    library.set_settings({"embedding_model": "fixture-v1", "embedding_status": "ready"})
    monkeypatch.setattr(Agent, "embedding", lambda *_: [1, 0])
    _, result = search(
        client, "conceptual coastal resilience", filters={"source_id": source_ids[0]}
    )
    assert len(result["results"]) == 1
    assert result["results"][0]["source_id"] == source_ids[0]
    assert "Related meaning" in result["results"][0]["signals"]
    library.set_settings({"embedding_model": "fixture-v2"})
    _, result = search(client, "conceptual coastal resilience")
    assert result["results"] == []


def test_disabled_history_keeps_followups_without_persisting_queries(workspace, tmp_path):
    client, library, worker = workspace
    folder = tmp_path / "docs"
    folder.mkdir()
    (folder / "migration.md").write_text("Migration release planning")
    connect(client, worker, folder)
    client.patch("/api/v1/settings", json={"history": False})
    session, result = search(client, "migration release")
    assert result["query"] == "migration release"
    _, followup = search(client, "only PDFs", session)
    assert "migration" in followup["effective_query"]
    assert followup["results"] == []
    with library.connect() as db:
        assert all(
            row[0] == "" and row[1] == ""
            for row in db.execute("SELECT query,effective_query FROM turns")
        )
        assert all(
            json.loads(row[0])["query"] == ""
            for row in db.execute("SELECT result_json FROM turns WHERE result_json IS NOT NULL")
        )
        assert all(row[0] == "" for row in db.execute("SELECT query FROM recommendation_events"))


def test_operational_audit_does_not_store_query_text(workspace, tmp_path):
    client, library, worker = workspace
    folder = tmp_path / "docs"
    folder.mkdir()
    (folder / "notes.txt").write_text("Private migration phrase")
    connect(client, worker, folder)
    _, result = search(client, "private migration phrase")
    with library.connect() as db:
        events = [dict(r) for r in db.execute("SELECT * FROM audit_events")]
    assert "private migration phrase" not in json.dumps(events).lower()


def test_restore_rejects_running_job(workspace):
    client, library, _ = workspace
    job = library.enqueue(kind="embeddings")
    with library.connect() as db:
        db.execute("UPDATE jobs SET status='running' WHERE id=?", (job,))
    assert (
        client.post(
            "/api/v1/restore", files={"archive": ("backup.zip", b"bad", "application/zip")}
        ).status_code
        == 409
    )


def test_worker_extracts_all_seven_formats_with_citations(workspace, tmp_path):
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from reportlab.pdfgen.canvas import Canvas

    client, _, worker = workspace
    folder = tmp_path / "formats"
    folder.mkdir()
    for extension in ("txt", "md", "rst"):
        (folder / f"notes.{extension}").write_text("Orchid release planning details")
    document = Document()
    document.add_paragraph("Orchid release document")
    document.save(folder / "notes.docx")
    workbook = Workbook()
    workbook.active.title = "Budget"
    workbook.active.append(["Orchid", 100])
    workbook.save(folder / "notes.xlsx")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, 1000000, 1000000).text = "Orchid presentation"
    presentation.save(folder / "notes.pptx")
    canvas = Canvas(str(folder / "notes.pdf"))
    canvas.drawString(72, 720, "Orchid research findings")
    canvas.save()
    connect(client, worker, folder)
    items = client.get("/api/v1/files").json()["files"]
    assert len(items) == 7
    for item in items:
        preview = client.get(f"/api/v1/files/{item['id']}/preview").json()
        assert any("Orchid" in s["text"] for s in preview["sections"])
        if item["extension"] == ".pdf":
            assert preview["sections"][0]["label"] == "Page 1"
        elif item["extension"] == ".pptx":
            assert preview["sections"][0]["label"] == "Slide 1"
        elif item["extension"] == ".xlsx":
            assert preview["sections"][0]["label"] == "Sheet: Budget"


def test_external_tracing_is_disabled_even_if_inherited(workspace, tmp_path, monkeypatch):
    from langsmith.run_helpers import get_tracing_context

    client, _, worker = workspace
    folder = tmp_path / "docs"
    folder.mkdir()
    (folder / "notes.txt").write_text("Local private project")
    connect(client, worker, folder)
    monkeypatch.setenv("LANGSMITH_TRACING", "true")
    graph = client.app.state.agent.graph
    original = graph.invoke
    observed = []

    def checked(*args, **kwargs):
        observed.append(get_tracing_context()["enabled"])
        return original(*args, **kwargs)

    monkeypatch.setattr(graph, "invoke", checked)
    search(client, "private project")
    assert observed == [False]


def test_search_recovery_happens_at_startup_not_factory_construction(tmp_path):
    from file_recommender.storage import LibraryDB, now

    store = IndexStore(tmp_path / "data/index.sqlite3")
    library = LibraryDB(store)
    with library.connect() as db:
        db.execute("INSERT INTO sessions(id,created_at) VALUES('s',?)", (now(),))
        db.execute(
            "INSERT INTO turns(id,session_id,query,status,created_at,updated_at) VALUES('t','s','query','running',?,?)",
            (now(), now()),
        )
    app = create_app(store)
    with library.connect() as db:
        assert db.execute("SELECT status FROM turns WHERE id='t'").fetchone()[0] == "running"
    with TestClient(app):
        with library.connect() as db:
            assert (
                db.execute("SELECT status FROM turns WHERE id='t'").fetchone()[0] == "interrupted"
            )
